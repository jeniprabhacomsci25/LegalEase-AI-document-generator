import os
import google.generativeai as genai
from docx import Document
from io import BytesIO
from dotenv import load_dotenv

# Load environment variables
load_dotenv()

def generate_legal_document(doc_type, parties_involved, terms_conditions, effective_date):
    """
    Generates a legal document using Google Gemini based on the updated UI parameters.
    """
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        return "Error: API Key is missing. Please enter a valid Gemini API Key."
        
    genai.configure(api_key=api_key)
    
    prompt = f"""
    Act as an expert corporate lawyer. Generate a comprehensive, legally binding {doc_type}. 
    
    Parties Involved:
    {parties_involved}
    
    Effective Date: {effective_date}
    
    Terms, Conditions, and Specifics: 
    {terms_conditions}
    
    Requirements:
    1. Structure the contract with clear, numbered sections.
    2. Use precise, professional legal terminology.
    3. Include necessary boilerplate clauses appropriate for this document type.
    4. Output the pure document text starting directly with the document title. Do not include introductory conversational text.
    """
    
    # 🔴 FIX: Disable Safety Filters for Legal Terminology
    safety_settings = [
        {"category": "HARM_CATEGORY_HARASSMENT", "threshold": "BLOCK_NONE"},
        {"category": "HARM_CATEGORY_HATE_SPEECH", "threshold": "BLOCK_NONE"},
        {"category": "HARM_CATEGORY_SEXUALLY_EXPLICIT", "threshold": "BLOCK_NONE"},
        {"category": "HARM_CATEGORY_DANGEROUS_CONTENT", "threshold": "BLOCK_NONE"}
    ]
    
    try:
        model = genai.GenerativeModel('gemini-3.5-flash') 
        # Pass the safety settings to the model
        response = model.generate_content(prompt, safety_settings=safety_settings)
        
        # Safely extract text
        try:
            return response.text
        except ValueError:
            # If Gemini still blocks it, show the exact reason instead of crashing
            reason = response.candidates[0].finish_reason.name if response.candidates else "Unknown"
            return f"Error: Gemini blocked the response due to safety filters. Reason: {reason}"
            
    except Exception as e:
        return f"Error generating document: {e}"

def create_docx(text, doc_title):
    """
    Converts the generated text into a formatted Word Document (.docx).
    """
    doc = Document()
    doc.add_heading(doc_title, 0)
    
    for paragraph in text.split('\n'):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())
            
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream
